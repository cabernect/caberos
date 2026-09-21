"""File preview renderers — bounded, format-classified payloads (W3).

Bytes in → JSON payload out. Pure functions: no DB, no filesystem access.
Bounds keep previews cheap on large files; binary formats (images, media,
PDF pages) are served through the raw-bytes endpoint so the panel streams
them via object URLs rather than inflating the JSON.

SVG is classified as `image` and served to an <img> tag — browsers never
execute scripts in image-loaded SVG, which is the sanitization strategy
(inline <svg> would need real XML sanitization and is intentionally not
offered).
"""

import base64
import csv
import io
import json
import mimetypes
from pathlib import Path

# Bounds — generous for real files, hard enough to keep panels snappy.
PREVIEW_MAX_BYTES = 64 * 1024 * 1024  # refuse to parse beyond this
TEXT_CHARS = 200_000
TABLE_MAX_ROWS = 500
TABLE_MAX_COLS = 64
XLSX_MAX_ROWS = 500
XLSX_MAX_COLS = 64
ELEMENTS_MAX = 400
ELEMENT_IMAGES_MAX = 20
ELEMENT_IMAGE_MAX_BYTES = 3 * 1024 * 1024
PDF_PAGE_SCALE = 1.5

_MARKDOWN = {".md", ".markdown", ".mdx"}
_CODE = {
    ".py": "python",
    ".js": "javascript",
    ".jsx": "jsx",
    ".ts": "typescript",
    ".tsx": "tsx",
    ".css": "css",
    ".html": "html",
    ".xml": "xml",
    ".yaml": "yaml",
    ".yml": "yaml",
    ".toml": "toml",
    ".sh": "bash",
    ".bash": "bash",
    ".zsh": "bash",
    ".sql": "sql",
    ".rs": "rust",
    ".go": "go",
    ".java": "java",
    ".c": "c",
    ".h": "c",
    ".cpp": "cpp",
    ".rb": "ruby",
    ".php": "php",
    ".swift": "swift",
    ".kt": "kotlin",
    ".lua": "lua",
    ".r": "r",
    ".vue": "vue",
    ".svelte": "svelte",
}
_TEXT = {".txt", ".log", ".env", ".gitignore", ".cfg", ".ini", ".conf"}
_JSON = {".json", ".jsonl", ".ndjson"}
_CSV = {".csv": ",", ".tsv": "\t"}
_IMAGE = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".ico", ".avif"}
_AUDIO = {".mp3", ".wav", ".ogg", ".m4a", ".flac", ".aac", ".opus"}
_VIDEO = {".mp4", ".webm", ".mov", ".mkv", ".avi", ".m4v"}

_MIME_OVERRIDES = {
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".svg": "image/svg+xml",
    ".js": "text/javascript",
    ".mjs": "text/javascript",
    ".ts": "text/plain",
    ".tsx": "text/plain",
    ".yaml": "text/yaml",
    ".yml": "text/yaml",
    ".toml": "text/plain",
    ".log": "text/plain",
    ".env": "text/plain",
}


def classify(name: str) -> str:
    """Filename → preview kind."""
    suffix = Path(name).suffix.lower()
    if suffix in _MARKDOWN:
        return "markdown"
    if suffix in _JSON:
        return "json"
    if suffix in _CSV:
        return "table"
    if suffix in _CODE:
        return "code"
    if suffix in _TEXT:
        return "text"
    if suffix in _IMAGE or suffix == ".svg":
        return "image"
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".docx":
        return "document"
    if suffix == ".pptx":
        return "slides"
    if suffix == ".xlsx":
        return "workbook"
    if suffix in _AUDIO or suffix in _VIDEO:
        return "media"
    return "unknown"


def media_type(name: str) -> str:
    """Content-Type for raw serving — overrides where mimetypes is wrong
    or absent for text-ish formats the panel renders itself."""
    suffix = Path(name).suffix.lower()
    if suffix in _MIME_OVERRIDES:
        return _MIME_OVERRIDES[suffix]
    return mimetypes.guess_type(name)[0] or "application/octet-stream"


def _media_kind(name: str) -> str:
    return "audio" if Path(name).suffix.lower() in _AUDIO else "video"


def preview_bytes(data: bytes, name: str) -> dict:
    """Bytes → bounded preview payload keyed by kind."""
    kind = classify(name)
    size = len(data)
    base = {"kind": kind, "size": size, "name": Path(name).name}
    if size > PREVIEW_MAX_BYTES:
        return {**base, "kind": "unknown", "too_large": True, "mime": media_type(name)}

    if kind in ("markdown", "code", "text"):
        text = data.decode("utf-8", errors="replace")
        payload = {**base, "content": text[:TEXT_CHARS], "truncated": len(text) > TEXT_CHARS}
        if kind == "code":
            payload["language"] = _CODE[Path(name).suffix.lower()]
        return payload
    if kind == "json":
        text = data.decode("utf-8", errors="replace")
        try:
            json.loads(text)
            valid = True
        except ValueError:
            valid = False
        return {
            **base,
            "content": text[:TEXT_CHARS],
            "truncated": len(text) > TEXT_CHARS,
            "valid": valid,
            "language": "json",
        }
    if kind == "table":
        return {**base, **_preview_table(data, name)}
    if kind == "image":
        return {**base, "mime": media_type(name), "svg": name.lower().endswith(".svg")}
    if kind == "pdf":
        return {**base, **_preview_pdf(data)}
    if kind == "document":
        return {**base, "format": "docx", **_preview_office(data, "docx")}
    if kind == "slides":
        return {**base, "format": "pptx", **_preview_slides(data)}
    if kind == "workbook":
        return {**base, "format": "xlsx", **_preview_workbook(data)}
    if kind == "media":
        return {**base, "media": _media_kind(name), "mime": media_type(name)}
    return {**base, "mime": media_type(name)}


def _preview_table(data: bytes, name: str) -> dict:
    """CSV/TSV → bounded header + rows. Sniffs the delimiter from the
    extension; reports the true row count even when truncated."""
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text), delimiter=_CSV[Path(name).suffix.lower()])
    rows: list[list[str]] = []
    total = 0
    for row in reader:
        total += 1
        if len(rows) < TABLE_MAX_ROWS:
            rows.append(row[:TABLE_MAX_COLS])
    truncated = total > len(rows) or any(len(r) >= TABLE_MAX_COLS for r in rows)
    header = rows[0] if rows else []
    return {
        "header": header,
        "rows": rows[1:],
        "total_rows": total,
        "truncated": truncated,
    }


def _preview_pdf(data: bytes) -> dict:
    """PDF → page count only; page PNGs render on demand via render_pdf_page."""
    try:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(data)
        count = len(pdf)
        pdf.close()
        return {"page_count": count}
    except Exception as e:
        return {"page_count": 0, "error": f"pdf unreadable: {e}"[:300]}


def render_pdf_page(data: bytes, page_number: int) -> bytes:
    """Render one PDF page to PNG bytes (bounded by PDF_PAGE_SCALE)."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    count = len(pdf)
    if page_number < 1 or page_number > count:
        pdf.close()
        raise ValueError(f"page {page_number} out of range (1-{count})")
    page = pdf[page_number - 1]
    bitmap = page.render(scale=PDF_PAGE_SCALE)
    image = bitmap.to_pil()
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=True)
    page.close()
    pdf.close()
    return buffer.getvalue()


def _preview_office(data: bytes, fmt: str) -> dict:
    """DOCX → normalized render elements (shared pdf_render vocabulary),
    images inlined as data URLs under caps."""
    try:
        from .artifacts import formats

        handler = formats.get_handler(fmt)
        elements = handler.to_elements(data, "")
    except Exception as e:
        return {"elements": [], "error": f"document unreadable: {e}"[:300]}
    capped, truncated = _cap_elements(elements)
    return {"elements": capped, "truncated": truncated}


def _preview_slides(data: bytes) -> dict:
    """PPTX → slides grouped on page_break markers from to_elements, plus
    the deck's canvas dims so the panel can frame slides at their true
    aspect ratio."""
    try:
        from .artifacts import formats

        elements = formats.get_handler("pptx").to_elements(data, "")
    except Exception as e:
        return {"slides": [], "error": f"presentation unreadable: {e}"[:300]}

    slide_width = slide_height = None
    try:
        from pptx import Presentation

        prs = Presentation(io.BytesIO(data))
        slide_width, slide_height = int(prs.slide_width), int(prs.slide_height)
        prs = None  # release the package refs promptly
    except Exception:
        pass  # dims are best-effort — the wireframe still renders without them

    slides: list[dict] = [{"title": None, "elements": []}]
    for el in elements:
        if el["type"] == "page_break":
            slides.append({"title": None, "elements": []})
            continue
        if el["type"] == "heading" and slides[-1]["title"] is None:
            slides[-1]["title"] = el.get("text")
        slides[-1]["elements"].append(el)

    import re

    generic_title = re.compile(r"^Slide \d+$")
    bare_bullet = re.compile(r"^[•·▪‣–—-]+$")
    for slide in slides:
        els = slide["elements"]
        # The injected leading heading is the title carrier — the card header
        # already displays it, so it must not render twice in the body.
        if els and els[0]["type"] == "heading" and els[0].get("level") == 1:
            els.pop(0)
        # Real decks often have no title placeholder (text boxes instead) —
        # fall back to the first short text run rather than "Slide N"; the
        # chosen element leaves the body so it never renders twice.
        if slide["title"] is None or generic_title.match(slide["title"] or ""):
            for j, e in enumerate(els):
                t = e.get("text", "").strip()
                if (
                    e["type"] == "paragraph"
                    and 0 < len(t) <= 80
                    and not bare_bullet.match(t)
                ):
                    slide["title"] = t
                    els.pop(j)
                    break
        # Bullet glyphs orphaned from their text by shape splitting, and
        # near-blank decoration images that render as empty white boxes.
        slide["elements"] = [
            e
            for e in els
            if not (
                e["type"] == "paragraph" and bare_bullet.match(e.get("text", "").strip())
            )
            and not (e["type"] == "image" and _near_blank_image(e.get("data")))
        ]
        # A body element repeating the resolved title verbatim (title
        # placeholders that also flow through the text path) renders as a
        # duplicate under the card header — drop it.
        title_text = (slide["title"] or "").strip()
        if title_text:
            slide["elements"] = [
                e
                for e in slide["elements"]
                if not (
                    e["type"] in ("heading", "paragraph")
                    and e.get("text", "").strip() == title_text
                )
            ]
        slide["elements"], slide["truncated"] = _cap_elements(slide["elements"])
    result = {"slides": slides, "slide_count": len(slides)}
    if slide_width and slide_height:
        result["slide_width"] = slide_width
        result["slide_height"] = slide_height
    return result


def _preview_workbook(data: bytes) -> dict:
    """XLSX → sheets with bounded row grids. Formula cells surface their
    formula string — openpyxl never evaluates, same honesty contract as
    the artifact inspect handler."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True)
    except Exception as e:
        return {"sheets": [], "error": f"workbook unreadable: {e}"[:300]}
    sheets = []
    for ws in wb.worksheets:
        # read_only worksheets may lack a declared dimension — fall back to
        # the bound (rstrip removes the padding it produces).
        max_col = min(ws.max_column, XLSX_MAX_COLS) if ws.max_column else XLSX_MAX_COLS
        rows: list[list] = []
        for row in ws.iter_rows(max_row=XLSX_MAX_ROWS, max_col=max_col):
            vals = [_json_safe(c.value) for c in row]
            while vals and vals[-1] is None:
                vals.pop()  # rstrip trailing empties — jagged rows, not padding
            rows.append(vals)
        while rows and not rows[-1]:
            rows.pop()  # drop empty tail rows
        sheets.append(
            {
                "name": ws.title,
                "rows": rows,
                "row_count": ws.max_row or 0,
                "truncated": (ws.max_row or 0) > XLSX_MAX_ROWS
                or (ws.max_column or 0) > XLSX_MAX_COLS,
            }
        )
    wb.close()
    return {"sheets": sheets, "formulas_recalculated": False}


def _json_safe(value):
    """openpyxl values → JSON-safe (datetimes → isoformat)."""
    import datetime

    if isinstance(value, datetime.datetime | datetime.date | datetime.time):
        return value.isoformat()
    return value


def _near_blank_image(data) -> bool:
    """Near-blank decoration images (white spacers, empty placeholder rects) —
    ≥97% pixels above ~245 luminance with a tiny encoded payload. Real
    graphics carry more data and detail, so both conditions must hold."""
    if not isinstance(data, bytes | bytearray) or len(data) > 8192:
        return False
    try:
        from PIL import Image

        im = Image.open(io.BytesIO(bytes(data))).convert("L")
        hist = im.histogram()
        return sum(hist[246:]) / (im.width * im.height) >= 0.97
    except Exception:
        return False


def _cap_elements(elements: list[dict]) -> tuple[list[dict], bool]:
    """Bound element count and inline image bytes as data URLs."""
    out: list[dict] = []
    images = 0
    truncated = False
    for el in elements[:ELEMENTS_MAX]:
        el = dict(el)
        if el["type"] == "image" and isinstance(el.get("data"), bytes | bytearray):
            if images >= ELEMENT_IMAGES_MAX or len(el["data"]) > ELEMENT_IMAGE_MAX_BYTES:
                truncated = True
                el = {"type": "notes", "text": "[image omitted — over preview limit]"}
            else:
                images += 1
                el = {"type": "image", "data_url": _data_url(bytes(el["data"]))}
        out.append(el)
    if len(elements) > ELEMENTS_MAX:
        truncated = True
        out.append({"type": "notes", "text": f"[{len(elements) - ELEMENTS_MAX} elements omitted]"})
    return out, truncated


def _data_url(data: bytes) -> str:
    mime = "image/png"
    if data[:3] == b"\xff\xd8\xff":
        mime = "image/jpeg"
    elif data[:6] in (b"GIF87a", b"GIF89a"):
        mime = "image/gif"
    elif data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        mime = "image/webp"
    return f"data:{mime};base64,{base64.b64encode(data).decode()}"

"""DOCX handler — python-docx structured build/revise/inspect/validate.

The model passes a block spec; this module owns the Office XML. Validation
reopens the file through zip + XML parsing independent of the writer path.
"""

import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree

from docx import Document
from docx.shared import Inches

from ...sandbox.workspace import resolve_within


def build(spec: dict, workspace_path: str | Path) -> bytes:
    doc = Document()
    if title := spec.get("title"):
        doc.add_heading(title, level=0)
    for block in spec.get("blocks", []):
        _add_block(doc, block, workspace_path)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def revise(data: bytes, ops: list[dict], workspace_path: str | Path) -> bytes:
    doc = Document(BytesIO(data))
    for op in ops:
        kind = op["op"]
        if kind == "append_blocks":
            for block in op["blocks"]:
                _add_block(doc, block, workspace_path)
        elif kind == "replace_paragraph":
            doc.paragraphs[op["index"]].text = op["text"]
        elif kind == "set_cell":
            t, r, c = op["table"], op["row"], op["col"]
            doc.tables[t].rows[r].cells[c].text = op["text"]
        else:
            raise ValueError(f"unknown docx op: {kind}")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def inspect(data: bytes) -> dict:
    doc = Document(BytesIO(data))
    headings, paragraphs, list_items = [], 0, 0
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            headings.append(p.text)
        elif p.style.name.startswith("List"):
            list_items += 1
        elif p.text.strip():
            paragraphs += 1
    return {
        "structure": {
            "headings": headings,
            "paragraphs": paragraphs,
            "list_items": list_items,
            "tables": len(doc.tables),
            "images": len(doc.inline_shapes),
        }
    }


def validate(data: bytes) -> dict:
    """Reopen via zip structure + XML parse — independent of the writer."""
    errors: list[str] = []
    try:
        zf = zipfile.ZipFile(BytesIO(data))
        if zf.testzip() is not None:
            errors.append("corrupt zip member")
        names = set(zf.namelist())
        for required in ("[Content_Types].xml", "word/document.xml"):
            if required not in names:
                errors.append(f"missing part: {required}")
        if "word/document.xml" in names:
            ElementTree.fromstring(zf.read("word/document.xml"))
    except zipfile.BadZipFile:
        errors.append("not a zip container")
    except ElementTree.ParseError as e:
        errors.append(f"document.xml not well-formed: {e}")
    if not errors:
        try:
            Document(BytesIO(data))
        except Exception as e:
            errors.append(f"parser reopen failed: {e}")
    return {"valid": not errors, "errors": errors}


def to_elements(data: bytes, workspace_path: str | Path) -> list[dict]:
    """Extract the normalized render elements (pdf_render vocabulary) from
    docx bytes — paragraphs/tables/images in document order."""
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    doc = Document(BytesIO(data))
    elements: list[dict] = []
    rels = doc.part.related_parts
    embed = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"

    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            p = DocxParagraph(child, doc)
            style = p.style.name if p.style is not None else ""
            if style == "Title" or style.startswith("Heading"):
                level = int(style[-1]) if style[-1:].isdigit() else 1
                elements.append({"type": "heading", "text": p.text, "level": level})
            elif style.startswith("List"):
                elements.append(
                    {
                        "type": "list",
                        "style": "number" if "Number" in style else "bullet",
                        "items": [p.text],
                    }
                )
            elif p.text.strip():
                elements.append({"type": "paragraph", "text": p.text})
            for blip in child.xpath('.//*[local-name()="blip"]'):
                rid = blip.get(embed)
                if rid and rid in rels:
                    elements.append({"type": "image", "data": rels[rid].blob})
        elif child.tag.endswith("}tbl"):
            t = DocxTable(child, doc)
            rows = [[c.text for c in r.cells] for r in t.rows]
            if rows:
                elements.append({"type": "table", "header": rows[0], "rows": rows[1:]})

    # Merge consecutive same-style lists so numbering continues correctly.
    merged: list[dict] = []
    for el in elements:
        if (
            el["type"] == "list"
            and merged
            and merged[-1]["type"] == "list"
            and merged[-1].get("style") == el.get("style")
        ):
            merged[-1]["items"].extend(el["items"])
        else:
            merged.append(el)
    return merged


def _add_block(doc, block: dict, workspace_path: str | Path) -> None:
    kind = block["type"]
    if kind == "heading":
        doc.add_heading(block["text"], level=block.get("level", 1))
    elif kind == "paragraph":
        doc.add_paragraph(block["text"])
    elif kind == "list":
        style = "List Number" if block.get("style") == "number" else "List Bullet"
        for item in block["items"]:
            doc.add_paragraph(str(item), style=style)
    elif kind == "table":
        rows = [block.get("header", []), *block.get("rows", [])]
        table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
        table.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, cell in enumerate(row):
                table.rows[i].cells[j].text = str(cell)
    elif kind == "image":
        img = resolve_within(workspace_path, block["path"])
        doc.add_picture(str(img), width=Inches(block.get("width_inches", 4)))
    elif kind == "page_break":
        doc.add_page_break()
    else:
        raise ValueError(f"unknown docx block: {kind}")

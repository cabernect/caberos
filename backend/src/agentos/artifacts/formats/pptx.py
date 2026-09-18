"""PPTX handler — python-pptx structured build/revise/inspect/validate."""

import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from ...sandbox.workspace import resolve_within

# Standard Office theme layout indices.
_LAYOUTS = {"title": 0, "title_content": 1, "title_only": 5, "blank": 6}
_CHARTS = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
}


def build(spec: dict, workspace_path: str | Path) -> bytes:
    prs = Presentation()
    for slide_spec in spec.get("slides", []):
        _add_slide(prs, slide_spec, workspace_path)
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def revise(data: bytes, ops: list[dict], workspace_path: str | Path) -> bytes:
    prs = Presentation(BytesIO(data))
    for op in ops:
        kind = op["op"]
        if kind == "add_slide":
            _add_slide(prs, op["slide"], workspace_path)
        elif kind == "move_slide":
            xml_slides = prs.slides._sldIdLst  # noqa: SLF001 — no public reorder API
            slides = list(xml_slides)
            xml_slides.remove(slides[op["from"]])
            xml_slides.insert(op["to"], slides[op["from"]])
        elif kind == "set_title":
            prs.slides[op["slide"]].shapes.title.text = op["text"]
        elif kind == "append_bullets":
            slide = prs.slides[op["slide"]]
            body = next((s for s in slide.placeholders if s.placeholder_format.idx == 1), None)
            if body is None:
                raise ValueError("slide has no content placeholder")
            for item in op["items"]:
                body.text_frame.add_paragraph().text = str(item)
        else:
            raise ValueError(f"unknown pptx op: {kind}")
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def inspect(data: bytes) -> dict:
    prs = Presentation(BytesIO(data))
    charts = tables = images = 0
    titles = []
    for slide in prs.slides:
        titles.append(slide.shapes.title.text if slide.shapes.title else None)
        for shape in slide.shapes:
            if shape.has_chart:
                charts += 1
            elif shape.has_table:
                tables += 1
            elif shape.shape_type == 13:  # PICTURE
                images += 1
    return {
        "structure": {
            "slides": len(prs.slides),
            "titles": titles,
            "charts": charts,
            "tables": tables,
            "images": images,
        }
    }


def validate(data: bytes) -> dict:
    errors: list[str] = []
    try:
        zf = zipfile.ZipFile(BytesIO(data))
        if zf.testzip() is not None:
            errors.append("corrupt zip member")
        names = set(zf.namelist())
        for required in ("[Content_Types].xml", "ppt/presentation.xml"):
            if required not in names:
                errors.append(f"missing part: {required}")
        if "ppt/presentation.xml" in names:
            ElementTree.fromstring(zf.read("ppt/presentation.xml"))
    except zipfile.BadZipFile:
        errors.append("not a zip container")
    except ElementTree.ParseError as e:
        errors.append(f"presentation.xml not well-formed: {e}")
    if not errors:
        try:
            Presentation(BytesIO(data))
        except Exception as e:
            errors.append(f"parser reopen failed: {e}")
    return {"valid": not errors, "errors": errors}


def _add_slide(prs: Presentation, slide_spec: dict, workspace_path: str | Path) -> None:
    layout = prs.slide_layouts[_LAYOUTS[slide_spec.get("layout", "title_content")]]
    slide = prs.slides.add_slide(layout)
    if title := slide_spec.get("title"):
        slide.shapes.title.text = title
    if subtitle := slide_spec.get("subtitle"):
        ph = next((s for s in slide.placeholders if s.placeholder_format.idx == 1), None)
        if ph is not None:
            ph.text = subtitle

    top = Inches(1.8)
    for block in slide_spec.get("blocks", []):
        kind = block["type"]
        if kind == "bullets":
            body = next((s for s in slide.placeholders if s.placeholder_format.idx == 1), None)
            if body is None:
                body = slide.shapes.add_textbox(Inches(0.7), top, Inches(8.6), Inches(4))
            tf = body.text_frame
            for i, item in enumerate(block["items"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = str(item)
        elif kind == "table":
            rows = [block.get("header", []), *block.get("rows", [])]
            shape = slide.shapes.add_table(
                len(rows),
                max(len(r) for r in rows),
                Inches(0.7),
                top,
                Inches(8.6),
                Inches(0.4 * len(rows)),
            )
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    shape.table.rows[i].cells[j].text = str(cell)
        elif kind == "image":
            img = resolve_within(workspace_path, block["path"])
            slide.shapes.add_picture(
                str(img), Inches(0.7), top, width=Inches(block.get("width_inches", 4))
            )
        elif kind == "chart":
            chart_data = CategoryChartData()
            chart_data.categories = block["categories"]
            for s in block["series"]:
                chart_data.add_series(s["name"], s["values"])
            slide.shapes.add_chart(
                _CHARTS[block.get("chart_type", "bar")],
                Inches(0.7),
                top,
                Inches(8.6),
                Inches(4.5),
                chart_data,
            )
        elif kind == "notes":
            slide.notes_slide.notes_text_frame.text = block["text"]
        elif kind == "text":
            box = slide.shapes.add_textbox(Inches(0.7), top, Inches(8.6), Inches(1))
            box.text_frame.text = block["text"]
        else:
            raise ValueError(f"unknown pptx block: {kind}")

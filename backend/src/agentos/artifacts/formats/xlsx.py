"""XLSX handler — openpyxl structured build/revise/inspect/validate.

Formula honesty: openpyxl persists formula strings without evaluating them.
`formulas_recalculated` is reported separately from authorship and is only
true when a real calculation engine (LibreOffice) produced values.
"""

import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook


def build(spec: dict, workspace_path: str | Path) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_spec in spec.get("sheets", [{"name": "Sheet1"}]):
        _build_sheet(wb, sheet_spec)
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def revise(data: bytes, ops: list[dict], workspace_path: str | Path) -> bytes:
    wb = load_workbook(BytesIO(data))
    for op in ops:
        kind = op["op"]
        if kind == "set_cell":
            _set_cell(wb[op["sheet"]], op["cell"], op)
        elif kind == "append_rows":
            for row in op["rows"]:
                wb[op["sheet"]].append(row)
        elif kind == "add_sheet":
            _build_sheet(wb, op)
        elif kind == "remove_sheet":
            wb.remove(wb[op["sheet"]])
        else:
            raise ValueError(f"unknown xlsx op: {kind}")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def inspect(data: bytes) -> dict:
    wb = load_workbook(BytesIO(data))
    formulas = 0
    cells = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cells += 1
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas += 1
    return {
        "structure": {
            "sheets": wb.sheetnames,
            "cells": cells,
            "formulas": formulas,
        },
        "formulas_recalculated": False,
    }


def validate(data: bytes) -> dict:
    errors: list[str] = []
    try:
        zf = zipfile.ZipFile(BytesIO(data))
        if zf.testzip() is not None:
            errors.append("corrupt zip member")
        names = set(zf.namelist())
        for required in ("[Content_Types].xml", "xl/workbook.xml"):
            if required not in names:
                errors.append(f"missing part: {required}")
        if "xl/workbook.xml" in names:
            ElementTree.fromstring(zf.read("xl/workbook.xml"))
    except zipfile.BadZipFile:
        errors.append("not a zip container")
    except ElementTree.ParseError as e:
        errors.append(f"workbook.xml not well-formed: {e}")
    if not errors:
        try:
            load_workbook(BytesIO(data))
        except Exception as e:
            errors.append(f"parser reopen failed: {e}")
    return {"valid": not errors, "errors": errors}


def to_elements(data: bytes, workspace_path: str | Path) -> list[dict]:
    """Extract render elements — one section per sheet, data as tables.
    Formula cells render as their formula string (unevaluated — same
    honesty contract as inspect's formulas_recalculated: false)."""
    wb = load_workbook(BytesIO(data))
    elements: list[dict] = []
    for i, ws in enumerate(wb.worksheets):
        if i:
            elements.append({"type": "page_break"})
        elements.append({"type": "heading", "text": ws.title, "level": 2})
        rows = []
        for row in ws.iter_rows():
            vals = ["" if c.value is None else str(c.value) for c in row]
            if any(vals):
                rows.append(vals)
        if rows:
            elements.append({"type": "table", "header": rows[0], "rows": rows[1:]})
        else:
            elements.append({"type": "paragraph", "text": "(empty sheet)"})
    return elements


def _build_sheet(wb: Workbook, sheet_spec: dict) -> None:
    ws = wb.create_sheet(sheet_spec.get("name", "Sheet"))
    for row in sheet_spec.get("rows", []):
        ws.append(row)
    for ref, cell_spec in sheet_spec.get("cells", {}).items():
        _set_cell(ws, ref, cell_spec)
    if freeze := sheet_spec.get("freeze"):
        ws.freeze_panes = freeze
    for col, width in sheet_spec.get("column_widths", {}).items():
        ws.column_dimensions[col].width = width
    for fmt_range, fmt in sheet_spec.get("formats", {}).items():
        for row in ws[fmt_range]:
            for cell in row:
                _apply_format(cell, fmt)


def _set_cell(ws, ref: str, cell_spec: dict | str | int | float) -> None:
    if not isinstance(cell_spec, dict):
        cell_spec = {"value": cell_spec}
    if "formula" in cell_spec:
        ws[ref] = cell_spec["formula"]
    elif "value" in cell_spec:
        ws[ref] = cell_spec["value"]
    _apply_format(ws[ref], cell_spec)


def _apply_format(cell, fmt: dict) -> None:
    from copy import copy

    if fmt.get("number_format"):
        cell.number_format = fmt["number_format"]
    font_changes = {k: fmt[k] for k in ("bold", "italic", "color") if k in fmt}
    if font_changes:
        font = copy(cell.font)
        for k, v in font_changes.items():
            setattr(font, k, v)
        cell.font = font

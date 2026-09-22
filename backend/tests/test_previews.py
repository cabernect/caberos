"""File previews — classify, bounded renderers, and the preview API (W3a).

Seam under test: agentos.previews renderers (pure bytes→payload) plus the
/api/agents/{id}/workspace/preview|raw|pdf-page endpoints. Tests observe
payloads and HTTP responses, never internals.
"""

from pathlib import Path

import pytest
from httpx import ASGITransport, AsyncClient

from agentos.db import get_db
from agentos.main import app


@pytest.fixture
async def client(db):
    """Test client with auth bypass — same pattern as the skills API tests."""
    from agentos.auth import require_operator
    from agentos.models.operator import Operator

    async def fake_operator():
        return Operator(id="test-operator", username="test", password_hash="x")

    app.dependency_overrides[require_operator] = fake_operator
    app.dependency_overrides[get_db] = lambda: db

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as c:
        yield c

    app.dependency_overrides.clear()


@pytest.fixture
def workspace_root(tmp_path, monkeypatch):
    """Point WorkspaceManager at tmp_path so workspaces are throwaway."""
    from agentos.config import settings

    root = tmp_path / "workspaces"
    root.mkdir()
    monkeypatch.setattr(settings, "workspace_root", root)
    return root


@pytest.fixture
def artifact_storage(tmp_path, monkeypatch):
    """Managed revision bytes under tmp_path, not the repo data dir."""
    from agentos.config import settings

    monkeypatch.setattr(settings, "db_path", tmp_path / "test.db")
    return tmp_path / "artifacts"


def _agent_workspace(root: Path, agent_id: str = "agent-1") -> Path:
    ws = root / agent_id
    ws.mkdir(parents=True, exist_ok=True)
    return ws


# --- classify + media_type ---


def test_classify_covers_plan_formats():
    from agentos import previews

    cases = {
        "notes.md": "markdown",
        "data.json": "json",
        "rows.csv": "table",
        "rows.tsv": "table",
        "script.py": "code",
        "readme.txt": "text",
        "shot.png": "image",
        "icon.svg": "image",
        "report.pdf": "pdf",
        "doc.docx": "document",
        "deck.pptx": "slides",
        "book.xlsx": "workbook",
        "song.mp3": "media",
        "clip.mp4": "media",
        "blob.bin": "unknown",
    }
    for name, kind in cases.items():
        assert previews.classify(name) == kind, name


def test_media_type_serves_svg_and_markdown_safely():
    from agentos import previews

    assert previews.media_type("icon.svg") == "image/svg+xml"
    assert previews.media_type("notes.md") == "text/markdown"
    assert previews.media_type("x.bin") == "application/octet-stream"


# --- bounded renderers ---


def test_text_preview_truncates():
    from agentos import previews

    payload = previews.preview_bytes(b"x" * (previews.TEXT_CHARS + 10), "big.txt")
    assert payload["kind"] == "text"
    assert payload["truncated"] is True
    assert len(payload["content"]) == previews.TEXT_CHARS


def test_json_preview_reports_validity():
    from agentos import previews

    assert previews.preview_bytes(b'{"a": 1}', "ok.json")["valid"] is True
    assert previews.preview_bytes(b"{nope", "bad.json")["valid"] is False


def test_code_preview_carries_language():
    from agentos import previews

    payload = previews.preview_bytes(b"print(1)", "tool.py")
    assert payload["kind"] == "code"
    assert payload["language"] == "python"


def test_csv_preview_bounds_rows_and_reports_total():
    from agentos import previews

    lines = ["a,b"] + [f"{i},{i}" for i in range(previews.TABLE_MAX_ROWS + 50)]
    payload = previews.preview_bytes("\n".join(lines).encode(), "big.csv")
    assert payload["kind"] == "table"
    assert payload["header"] == ["a", "b"]
    assert payload["total_rows"] == previews.TABLE_MAX_ROWS + 51
    assert len(payload["rows"]) == previews.TABLE_MAX_ROWS - 1
    assert payload["truncated"] is True


def test_image_preview_is_metadata_only():
    from agentos import previews

    payload = previews.preview_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 100, "p.png")
    assert payload["kind"] == "image"
    assert payload["mime"] == "image/png"
    assert "content" not in payload  # bytes stream through /raw, not JSON


def test_svg_preview_uses_image_kind():
    """SVG → image kind: <img> loading is the script-execution sanitizer."""
    from agentos import previews

    payload = previews.preview_bytes(b"<svg><script>x</script></svg>", "i.svg")
    assert payload["kind"] == "image"
    assert payload["svg"] is True


def _pdf_bytes() -> bytes:
    from agentos.artifacts.pdf_render import render_document

    return render_document([{"type": "paragraph", "text": "hello pdf"}], ".")


def test_pdf_preview_reports_page_count():
    from agentos import previews

    payload = previews.preview_bytes(_pdf_bytes(), "doc.pdf")
    assert payload["kind"] == "pdf"
    assert payload["page_count"] == 1


def test_render_pdf_page_returns_png():
    from agentos import previews

    png = previews.render_pdf_page(_pdf_bytes(), 1)
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
    with pytest.raises(ValueError):
        previews.render_pdf_page(_pdf_bytes(), 99)


def test_docx_preview_extracts_elements():
    from agentos import previews
    from agentos.artifacts.formats import docx

    data = docx.build(
        {
            "title": "Report",
            "blocks": [
                {"type": "heading", "text": "Summary", "level": 1},
                {"type": "paragraph", "text": "body"},
                {"type": "table", "header": ["K"], "rows": [["v"]]},
            ],
        },
        ".",
    )
    payload = previews.preview_bytes(data, "doc.docx")
    assert payload["kind"] == "document"
    types = [e["type"] for e in payload["elements"]]
    assert "heading" in types and "paragraph" in types and "table" in types


def test_pptx_preview_groups_slides():
    from agentos import previews
    from agentos.artifacts.formats import pptx

    data = pptx.build(
        {
            "slides": [
                {"layout": "title", "title": "One"},
                {
                    "layout": "title_content",
                    "title": "Two",
                    "blocks": [{"type": "bullets", "items": ["a"]}],
                },
            ]
        },
        ".",
    )
    payload = previews.preview_bytes(data, "deck.pptx")
    assert payload["kind"] == "slides"
    assert payload["slide_count"] == 2
    assert payload["slides"][0]["title"] == "One"
    assert payload["slides"][1]["title"] == "Two"


def test_pptx_preview_untitled_slide_falls_back_to_text():
    """Real decks often lack title placeholders — the card title falls back
    to the first short text run, and the injected heading must not render
    twice in the body."""
    from agentos import previews
    from agentos.artifacts.formats import pptx

    data = pptx.build(
        {"slides": [{"layout": "title_only", "blocks": []}]},
        ".",
    )
    # title_only has a title placeholder; drop it to simulate a text-box deck.
    from io import BytesIO

    from pptx import Presentation

    prs = Presentation(BytesIO(data))
    slide = prs.slides[0]
    slide.shapes.title.text = ""
    buf = BytesIO()
    prs.save(buf)

    payload = previews.preview_bytes(buf.getvalue(), "deck.pptx")
    assert payload["kind"] == "slides"
    s1 = payload["slides"][0]
    # No paragraph text → generic fallback stays, but body must not carry a
    # duplicated injected heading.
    headings = [e for e in s1["elements"] if e["type"] == "heading"]
    assert headings == []
    assert s1["title"] is not None


def test_xlsx_preview_returns_sheet_grids():
    from agentos import previews
    from agentos.artifacts.formats import xlsx

    data = xlsx.build(
        {
            "sheets": [
                {
                    "name": "Rev",
                    "rows": [["Region", "Q3"], ["VN", 120]],
                    "cells": {"B3": {"formula": "=SUM(B2:B2)"}},
                }
            ]
        },
        ".",
    )
    payload = previews.preview_bytes(data, "book.xlsx")
    assert payload["kind"] == "workbook"
    sheet = payload["sheets"][0]
    assert sheet["name"] == "Rev"
    assert sheet["rows"][0] == ["Region", "Q3"]
    # Formula surfaces as its string — never silently evaluated.
    assert sheet["rows"][2][1] == "=SUM(B2:B2)"
    assert payload["formulas_recalculated"] is False


def test_oversized_file_returns_too_large():
    from agentos import previews

    payload = previews.preview_bytes(b"0" * (previews.PREVIEW_MAX_BYTES + 1), "huge.docx")
    assert payload["kind"] == "unknown"
    assert payload["too_large"] is True


# --- API ---


async def test_preview_workspace_text_file(client, workspace_root, artifact_storage):
    ws = _agent_workspace(workspace_root)
    (ws / "notes.md").write_text("# Hello\n\nworld")

    resp = await client.get("/api/agents/agent-1/workspace/preview?path=notes.md")
    assert resp.status_code == 200
    data = resp.json()
    assert data["kind"] == "markdown"
    assert "Hello" in data["content"]
    assert data["artifact"] is None


async def test_preview_containment_blocks_escape(client, workspace_root, artifact_storage):
    _agent_workspace(workspace_root)
    resp = await client.get("/api/agents/agent-1/workspace/preview?path=../secret.txt")
    assert resp.status_code in (403, 404)


async def test_preview_tracked_file_embeds_artifact_meta(
    db, client, workspace_root, artifact_storage
):
    from agentos.artifacts import service

    ws = _agent_workspace(workspace_root)
    art, rev = await service.create(
        db,
        workspace_id="agent-1",
        workspace_path=ws,
        rel_path="report.docx",
        data=b"docx-v1",
    )
    resp = await client.get("/api/agents/agent-1/workspace/preview?path=report.docx")
    assert resp.status_code == 200
    meta = resp.json()["artifact"]
    assert meta["id"] == art.id
    assert meta["revision_count"] == 1
    assert meta["current_revision_number"] == 1
    assert meta["newer_exists"] is False


async def test_preview_revision_uses_stored_bytes_and_flags_newer(
    db, client, workspace_root, artifact_storage
):
    """Chat previews the exact revision attached to a message — stored bytes,
    not the current file — and newer_exists drives the banner."""
    from agentos.artifacts import service

    ws = _agent_workspace(workspace_root)
    art, rev1 = await service.create(
        db,
        workspace_id="agent-1",
        workspace_path=ws,
        rel_path="doc.txt",
        data=b"version-one",
    )
    rev2 = await service.revise(
        db, art.id, workspace_path=ws, base_revision_id=rev1.id, data=b"version-two"
    )

    resp = await client.get(
        f"/api/agents/agent-1/workspace/preview?artifact_id={art.id}&revision_id={rev1.id}"
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["content"] == "version-one"  # old bytes, not the current file
    meta = data["artifact"]
    assert meta["viewing_revision_number"] == 1
    assert meta["current_revision_number"] == 2
    assert meta["newer_exists"] is True

    # Current revision shows no banner.
    resp = await client.get(
        f"/api/agents/agent-1/workspace/preview?artifact_id={art.id}&revision_id={rev2.id}"
    )
    assert resp.json()["artifact"]["newer_exists"] is False


async def test_preview_rejects_cross_agent_artifact(db, client, workspace_root, artifact_storage):
    from agentos.artifacts import service

    ws = _agent_workspace(workspace_root)
    art, _ = await service.create(
        db,
        workspace_id="agent-1",
        workspace_path=ws,
        rel_path="doc.txt",
        data=b"x",
    )
    resp = await client.get(f"/api/agents/other-agent/workspace/preview?artifact_id={art.id}")
    assert resp.status_code == 404


async def test_raw_serves_bytes_with_content_type(client, workspace_root, artifact_storage):
    ws = _agent_workspace(workspace_root)
    (ws / "img.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 50)

    resp = await client.get("/api/agents/agent-1/workspace/raw?path=img.png")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("image/png")
    assert resp.content[:8] == b"\x89PNG\r\n\x1a\n"


async def test_raw_serves_revision_bytes(db, client, workspace_root, artifact_storage):
    from agentos.artifacts import service

    ws = _agent_workspace(workspace_root)
    art, rev1 = await service.create(
        db,
        workspace_id="agent-1",
        workspace_path=ws,
        rel_path="a.txt",
        data=b"old",
    )
    await service.revise(db, art.id, workspace_path=ws, base_revision_id=rev1.id, data=b"new")
    resp = await client.get(
        f"/api/agents/agent-1/workspace/raw?artifact_id={art.id}&revision_id={rev1.id}"
    )
    assert resp.content == b"old"


async def test_pdf_page_endpoint_renders_png(client, workspace_root, artifact_storage):
    ws = _agent_workspace(workspace_root)
    (ws / "doc.pdf").write_bytes(_pdf_bytes())

    resp = await client.get("/api/agents/agent-1/workspace/pdf-page?path=doc.pdf&page=1")
    assert resp.status_code == 200
    assert resp.content[:8] == b"\x89PNG\r\n\x1a\n"
    assert resp.headers["content-type"] == "image/png"

    resp = await client.get("/api/agents/agent-1/workspace/pdf-page?path=doc.pdf&page=9")
    assert resp.status_code == 404


async def test_pdf_page_rejects_non_pdf(client, workspace_root, artifact_storage):
    ws = _agent_workspace(workspace_root)
    (ws / "a.txt").write_text("not a pdf")
    resp = await client.get("/api/agents/agent-1/workspace/pdf-page?path=a.txt&page=1")
    assert resp.status_code == 400


# --- Artifact panel actions ---


async def test_revisions_endpoint_lists_newest_first(db, client, workspace_root, artifact_storage):
    from agentos.artifacts import service

    ws = _agent_workspace(workspace_root)
    art, rev1 = await service.create(
        db, workspace_id="agent-1", workspace_path=ws, rel_path="a.txt", data=b"1"
    )
    await service.revise(db, art.id, workspace_path=ws, base_revision_id=rev1.id, data=b"2")
    resp = await client.get(f"/api/agents/agent-1/artifacts/{art.id}/revisions")
    assert resp.status_code == 200
    revs = resp.json()["revisions"]
    assert [r["revision_number"] for r in revs] == [2, 1]
    assert revs[0]["current"] is True

    resp = await client.get(f"/api/agents/other/artifacts/{art.id}/revisions")
    assert resp.status_code == 404


async def test_restore_endpoint_appends_new_revision(db, client, workspace_root, artifact_storage):
    from agentos.artifacts import service

    ws = _agent_workspace(workspace_root)
    art, rev1 = await service.create(
        db, workspace_id="agent-1", workspace_path=ws, rel_path="a.txt", data=b"old"
    )
    await service.revise(db, art.id, workspace_path=ws, base_revision_id=rev1.id, data=b"new")
    resp = await client.post(
        f"/api/agents/agent-1/artifacts/{art.id}/restore",
        json={"revision_id": rev1.id},
    )
    assert resp.status_code == 200
    assert resp.json()["revision_number"] == 3
    assert (ws / "a.txt").read_bytes() == b"old"

    resp = await client.post(
        f"/api/agents/other/artifacts/{art.id}/restore",
        json={"revision_id": rev1.id},
    )
    assert resp.status_code == 404


async def test_adopt_endpoint_tracks_ordinary_file(db, client, workspace_root, artifact_storage):
    ws = _agent_workspace(workspace_root)
    (ws / "drop.txt").write_text("user file")

    resp = await client.post("/api/agents/agent-1/artifacts/adopt", json={"path": "drop.txt"})
    assert resp.status_code == 200
    assert resp.json()["path"] == "drop.txt"

    # Now the preview embeds artifact metadata.
    resp = await client.get("/api/agents/agent-1/workspace/preview?path=drop.txt")
    assert resp.json()["artifact"]["revision_count"] == 1

    resp = await client.post("/api/agents/agent-1/artifacts/adopt", json={"path": "missing.txt"})
    assert resp.status_code == 400


async def test_compare_endpoint_diffs_text_revisions(db, client, workspace_root, artifact_storage):
    from agentos.artifacts import service

    ws = _agent_workspace(workspace_root)
    art, rev1 = await service.create(
        db,
        workspace_id="agent-1",
        workspace_path=ws,
        rel_path="notes.txt",
        data=b"line one\nline two\n",
    )
    rev2 = await service.revise(
        db,
        art.id,
        workspace_path=ws,
        base_revision_id=rev1.id,
        data=b"line one\nline changed\nline three\n",
    )
    resp = await client.get(
        f"/api/agents/agent-1/artifacts/{art.id}/compare?from_revision={rev1.id}"
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["comparable"] is True
    assert data["from_revision_number"] == 1
    assert data["to_revision_number"] == 2
    assert "-line two" in data["diff"]
    assert "+line changed" in data["diff"]
    assert "+line three" in data["diff"]

    # Explicit pair + cross-agent isolation.
    resp = await client.get(
        f"/api/agents/agent-1/artifacts/{art.id}/compare"
        f"?from_revision={rev1.id}&to_revision={rev2.id}"
    )
    assert resp.status_code == 200
    resp = await client.get(f"/api/agents/other/artifacts/{art.id}/compare?from_revision={rev1.id}")
    assert resp.status_code == 404


async def test_compare_endpoint_reports_binary_honestly(
    db, client, workspace_root, artifact_storage
):
    from agentos.artifacts import service

    ws = _agent_workspace(workspace_root)
    art, rev1 = await service.create(
        db,
        workspace_id="agent-1",
        workspace_path=ws,
        rel_path="blob.bin",
        data=b"\x00\xff\x01",
    )
    await service.revise(
        db, art.id, workspace_path=ws, base_revision_id=rev1.id, data=b"\x00\xff\x02\x03"
    )
    resp = await client.get(
        f"/api/agents/agent-1/artifacts/{art.id}/compare?from_revision={rev1.id}"
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["comparable"] is False
    assert data["reason"] == "binary"
    assert data["identical"] is False
    assert data["from_bytes"] == 3 and data["to_bytes"] == 4


# --- Composer attachment previews (W3c) ---


async def test_ephemeral_preview_renders_without_persisting(client, workspace_root):
    ws = _agent_workspace(workspace_root)
    resp = await client.post(
        "/api/agents/agent-1/attachments/preview",
        files={"file": ("table.csv", b"a,b\n1,2\n3,4\n", "text/csv")},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["kind"] == "table"
    assert data["rows"][0] == ["1", "2"]
    # Nothing landed in the workspace — preview is ephemeral.
    assert list(ws.iterdir()) == []


async def test_ephemeral_preview_pdf_carries_thumbnail(client, workspace_root):
    _agent_workspace(workspace_root)
    resp = await client.post(
        "/api/agents/agent-1/attachments/preview",
        files={"file": ("doc.pdf", _pdf_bytes(), "application/pdf")},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["kind"] == "pdf"
    import base64

    assert base64.b64decode(data["thumb_png"])[:8] == b"\x89PNG\r\n\x1a\n"


async def test_url_preview_requires_http(client, workspace_root):
    _agent_workspace(workspace_root)
    resp = await client.get("/api/agents/agent-1/attachments/url-preview?url=file:///etc/passwd")
    assert resp.status_code == 400


async def test_url_preview_extracts_title(client, workspace_root, monkeypatch):
    """Title extraction is best-effort — a mocked fetch returns the page."""
    _agent_workspace(workspace_root)

    class FakeResp:
        content = b"<html><head><title>  Example   Page </title></head></html>"

    class FakeClient:
        def __init__(self, **kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *a):
            return False

        async def get(self, url):
            return FakeResp()

    import httpx

    monkeypatch.setattr(httpx, "AsyncClient", FakeClient)
    resp = await client.get(
        "/api/agents/agent-1/attachments/url-preview?url=https://example.com/page"
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["domain"] == "example.com"
    assert data["title"] == "Example Page"


# --- Skill resource previews (W3b) ---
#
# Same renderers, different root: skill dirs are operator-managed content
# outside any agent workspace, so containment is anchored at the skill dir.


@pytest.fixture
def skills_dir(tmp_path, monkeypatch):
    """Point the skills API at a throwaway skill directory."""
    from agentos.api import skills as skills_api

    root = tmp_path / "skills"
    (root / "demo-skill" / "refs").mkdir(parents=True)
    (root / "demo-skill" / "SKILL.md").write_text("---\nname: demo-skill\n---\n# Demo")
    (root / "demo-skill" / "refs" / "guide.md").write_text("# Guide\n\nsteps")
    (root / "demo-skill" / "refs" / "data.csv").write_text("a,b\n1,2\n")
    monkeypatch.setattr(skills_api, "SKILLS_DIR", root)
    return root


async def test_skill_resources_lists_files(client, skills_dir):
    resp = await client.get("/api/skills/demo-skill/resources")
    assert resp.status_code == 200
    paths = {r["path"] for r in resp.json()["resources"]}
    assert "SKILL.md" in paths
    assert "refs/guide.md" in paths

    resp = await client.get("/api/skills/no-such/resources")
    assert resp.status_code == 404


async def test_skill_preview_renders_markdown(client, skills_dir):
    resp = await client.get("/api/skills/demo-skill/preview?path=refs/guide.md")
    assert resp.status_code == 200
    data = resp.json()
    assert data["kind"] == "markdown"
    assert "Guide" in data["content"]
    assert data["artifact"] is None


async def test_skill_preview_containment_blocks_escape(client, skills_dir):
    resp = await client.get("/api/skills/demo-skill/preview?path=../outside.txt")
    assert resp.status_code in (403, 404)
    resp = await client.get("/api/skills/..%2fpreview?path=x")
    assert resp.status_code in (400, 404, 422)


async def test_skill_raw_and_pdf_page(client, skills_dir):
    (skills_dir / "demo-skill" / "refs" / "doc.pdf").write_bytes(_pdf_bytes())

    resp = await client.get("/api/skills/demo-skill/raw?path=refs/data.csv")
    assert resp.status_code == 200
    assert resp.content == b"a,b\n1,2\n"

    resp = await client.get("/api/skills/demo-skill/pdf-page?path=refs/doc.pdf&page=1")
    assert resp.status_code == 200
    assert resp.content[:8] == b"\x89PNG\r\n\x1a\n"

    resp = await client.get("/api/skills/demo-skill/pdf-page?path=refs/data.csv&page=1")
    assert resp.status_code == 400


def test_pptx_preview_reports_slide_dims():
    """The payload carries canvas dims so the panel can frame the true aspect."""
    from agentos import previews
    from agentos.artifacts.formats import pptx

    data = pptx.build({"slides": [{"layout": "title", "title": "T"}]}, ".")
    payload = previews.preview_bytes(data, "deck.pptx")
    assert payload["slide_width"] == 12192000
    assert payload["slide_height"] == 6858000


def test_pptx_preview_dedupes_title_text():
    """A body element repeating the resolved title verbatim is dropped —
    the card header already shows it."""
    from agentos import previews
    from agentos.artifacts.formats import pptx

    data = pptx.build(
        {
            "slides": [
                {
                    "layout": "title_content",
                    "title": "Same",
                    "blocks": [{"type": "text", "text": "Same"}],
                }
            ]
        },
        ".",
    )
    payload = previews.preview_bytes(data, "deck.pptx")
    s1 = payload["slides"][0]
    assert s1["title"] == "Same"
    assert not any(
        e["type"] == "paragraph" and e.get("text", "").strip() == "Same" for e in s1["elements"]
    )


# --- Workspace delete (operator file management) ---


async def test_workspace_delete_file(client, workspace_root):
    ws = _agent_workspace(workspace_root)
    (ws / "old.txt").write_text("gone")

    resp = await client.delete("/api/agents/agent-1/workspace?path=old.txt")
    assert resp.status_code == 200
    assert resp.json()["deleted"] is True
    assert not (ws / "old.txt").exists()


async def test_workspace_delete_dir_recursive(client, workspace_root):
    ws = _agent_workspace(workspace_root)
    (ws / "scratch" / "sub").mkdir(parents=True)
    (ws / "scratch" / "a.txt").write_text("a")
    (ws / "scratch" / "sub" / "b.txt").write_text("b")

    resp = await client.delete("/api/agents/agent-1/workspace?path=scratch")
    assert resp.status_code == 200
    assert not (ws / "scratch").exists()


async def test_workspace_delete_missing_returns_404(client, workspace_root):
    _agent_workspace(workspace_root)
    resp = await client.delete("/api/agents/agent-1/workspace?path=nope.txt")
    assert resp.status_code == 404


async def test_workspace_delete_requires_path(client, workspace_root):
    _agent_workspace(workspace_root)
    resp = await client.delete("/api/agents/agent-1/workspace")
    assert resp.status_code == 400


async def test_workspace_delete_rejects_escape(client, workspace_root):
    _agent_workspace(workspace_root)
    resp = await client.delete("/api/agents/agent-1/workspace?path=../outside.txt")
    assert resp.status_code == 403


async def test_workspace_delete_refuses_tracked_artifact(client, workspace_root, db):
    from agentos.models.artifact import Artifact

    ws = _agent_workspace(workspace_root)
    (ws / "artifacts").mkdir()
    (ws / "artifacts" / "deck.pptx").write_bytes(b"pptx-bytes")
    db.add(Artifact(workspace_id="agent-1", current_path="artifacts/deck.pptx", format="pptx"))
    await db.commit()

    resp = await client.delete("/api/agents/agent-1/workspace?path=artifacts/deck.pptx")
    assert resp.status_code == 409
    assert "tracked" in resp.json()["detail"]
    assert (ws / "artifacts" / "deck.pptx").exists()


async def test_workspace_delete_dir_refuses_when_artifact_inside(client, workspace_root, db):
    from agentos.models.artifact import Artifact

    ws = _agent_workspace(workspace_root)
    (ws / "artifacts").mkdir()
    (ws / "artifacts" / "deck.pptx").write_bytes(b"pptx-bytes")
    db.add(Artifact(workspace_id="agent-1", current_path="artifacts/deck.pptx", format="pptx"))
    await db.commit()

    resp = await client.delete("/api/agents/agent-1/workspace?path=artifacts")
    assert resp.status_code == 409
    assert (ws / "artifacts" / "deck.pptx").exists()


async def test_workspace_delete_allows_untracked_sibling(client, workspace_root, db):
    """An untracked file next to a tracked artifact still deletes."""
    from agentos.models.artifact import Artifact

    ws = _agent_workspace(workspace_root)
    (ws / "artifacts").mkdir()
    (ws / "artifacts" / "deck.pptx").write_bytes(b"pptx-bytes")
    (ws / "artifacts" / "notes.txt").write_text("loose")
    db.add(Artifact(workspace_id="agent-1", current_path="artifacts/deck.pptx", format="pptx"))
    await db.commit()

    resp = await client.delete("/api/agents/agent-1/workspace?path=artifacts/notes.txt")
    assert resp.status_code == 200
    assert not (ws / "artifacts" / "notes.txt").exists()
    assert (ws / "artifacts" / "deck.pptx").exists()


async def test_workspace_delete_symlink_leaf_unlinks_not_follows(client, workspace_root):
    """A symlink whose target escapes the workspace is itself a workspace
    entry — deleting must unlink the link, never follow it (A20)."""
    outside = workspace_root / "outside-secret.txt"
    outside.write_text("secret")
    ws = _agent_workspace(workspace_root)
    link = ws / "escape-link"
    link.symlink_to(outside)

    resp = await client.delete("/api/agents/agent-1/workspace?path=escape-link")
    assert resp.status_code == 200
    assert not link.exists() and not link.is_symlink()
    assert outside.exists()  # target untouched


async def test_workspace_delete_symlink_dir_leaf_unlinks(client, workspace_root):
    """A symlinked directory leaf is unlinked — rmtree must never follow it."""
    outside_dir = workspace_root / "outside-dir"
    outside_dir.mkdir()
    (outside_dir / "keep.txt").write_text("keep")
    ws = _agent_workspace(workspace_root)
    (ws / "dir-link").symlink_to(outside_dir, target_is_directory=True)

    resp = await client.delete("/api/agents/agent-1/workspace?path=dir-link")
    assert resp.status_code == 200
    assert (outside_dir / "keep.txt").exists()


async def test_workspace_delete_through_symlinked_dir_still_blocked(client, workspace_root):
    """Containment still applies through intermediate components — a file
    reached via a symlinked directory escapes and must refuse."""
    outside_dir = workspace_root / "outside-dir"
    outside_dir.mkdir()
    (outside_dir / "victim.txt").write_text("victim")
    ws = _agent_workspace(workspace_root)
    (ws / "dir-link").symlink_to(outside_dir, target_is_directory=True)

    resp = await client.delete("/api/agents/agent-1/workspace?path=dir-link/victim.txt")
    assert resp.status_code == 403
    assert (outside_dir / "victim.txt").exists()


async def test_workspace_delete_normalizes_dotdot(client, workspace_root):
    ws = _agent_workspace(workspace_root)
    (ws / "sub").mkdir()
    (ws / "target.txt").write_text("x")

    resp = await client.delete("/api/agents/agent-1/workspace?path=sub/../target.txt")
    assert resp.status_code == 200
    assert not (ws / "target.txt").exists()

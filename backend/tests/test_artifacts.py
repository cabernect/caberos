"""Artifact service — format-agnostic revision core (W2a).

Seam under test: agentos.artifacts.service — the module interface the
artifact capabilities call. Tests observe behavior through service calls +
workspace files, never internals.
"""

import hashlib
from pathlib import Path

import pytest


@pytest.fixture
def artifact_storage(tmp_path, monkeypatch):
    """Managed revision bytes under tmp_path, not the repo data dir."""
    from agentos.config import settings

    monkeypatch.setattr(settings, "db_path", tmp_path / "test.db")
    return tmp_path / "artifacts"


async def test_create_writes_workspace_file_and_first_revision(db, workspace, artifact_storage):
    from agentos.artifacts import service

    art, rev = await service.create(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="report.docx",
        data=b"docx-bytes-v1",
        created_by="agent-1",
    )

    assert (Path(workspace) / "report.docx").read_bytes() == b"docx-bytes-v1"
    assert art.format == "docx"
    assert art.tracking_status == "tracked"
    assert rev.revision_number == 1
    assert rev.content_hash == hashlib.sha256(b"docx-bytes-v1").hexdigest()
    assert art.current_revision_id == rev.id
    # Revision bytes preserved in managed storage, readable via the service.
    assert await service.revision_bytes(db, rev.id) == b"docx-bytes-v1"


async def test_revise_creates_new_revision_and_updates_file(db, workspace, artifact_storage):
    from agentos.artifacts import service

    art, rev1 = await service.create(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="deck.pptx",
        data=b"v1",
        created_by="agent-1",
    )

    rev2 = await service.revise(
        db,
        art.id,
        workspace_path=workspace,
        base_revision_id=rev1.id,
        data=b"v2",
        change_summary="edit title slide",
    )

    assert rev2.revision_number == 2
    assert rev2.base_revision_id == rev1.id
    assert (Path(workspace) / "deck.pptx").read_bytes() == b"v2"
    assert await service.revision_bytes(db, rev1.id) == b"v1"  # history preserved
    await db.refresh(art)
    assert art.current_revision_id == rev2.id


async def test_external_modification_conflicts_instead_of_overwrite(
    db, workspace, artifact_storage
):
    from agentos.artifacts import service

    art, rev1 = await service.create(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="book.xlsx",
        data=b"v1",
        created_by="agent-1",
    )
    # User edits the file outside the agent (e.g. in Excel).
    (Path(workspace) / "book.xlsx").write_bytes(b"external-edit")

    with pytest.raises(service.ArtifactConflictError):
        await service.revise(
            db,
            art.id,
            workspace_path=workspace,
            base_revision_id=rev1.id,
            data=b"agent-draft",
        )

    # The external edit is never overwritten, and no phantom revision exists.
    assert (Path(workspace) / "book.xlsx").read_bytes() == b"external-edit"
    await db.refresh(art)
    assert art.tracking_status == "conflict"
    assert art.current_revision_id == rev1.id


async def test_capture_external_stores_the_external_edit_as_a_revision(
    db, workspace, artifact_storage
):
    from agentos.artifacts import service

    art, rev1 = await service.create(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="book.xlsx",
        data=b"v1",
        created_by="agent-1",
    )
    (Path(workspace) / "book.xlsx").write_bytes(b"external-edit")

    ext = await service.capture_external(
        db, art.id, workspace_path=workspace, change_summary="user edit in Excel"
    )

    assert ext.revision_number == 2
    await db.refresh(art)
    assert art.tracking_status == "tracked"
    assert art.current_revision_id == ext.id
    # Agent can now revise on top of the captured external revision.
    rev3 = await service.revise(
        db,
        art.id,
        workspace_path=workspace,
        base_revision_id=ext.id,
        data=b"agent-merge",
    )
    assert rev3.revision_number == 3


async def test_restore_creates_a_new_revision(db, workspace, artifact_storage):
    """Restoring v2 after v5 produces v6 — history is never rewritten."""
    from agentos.artifacts import service

    art, rev1 = await service.create(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="doc.docx",
        data=b"v1",
        created_by="agent-1",
    )
    last = rev1
    for i in range(2, 6):
        last = await service.revise(
            db,
            art.id,
            workspace_path=workspace,
            base_revision_id=last.id,
            data=f"v{i}".encode(),
        )
    assert last.revision_number == 5

    rev6 = await service.restore(db, art.id, workspace_path=workspace, revision_id=rev1.id)

    assert rev6.revision_number == 6
    assert rev6.content_hash == rev1.content_hash
    assert rev6.base_revision_id == last.id
    assert (Path(workspace) / "doc.docx").read_bytes() == b"v1"
    await db.refresh(art)
    assert art.current_revision_id == rev6.id


async def test_paths_cannot_escape_the_workspace(db, workspace, artifact_storage):
    from agentos.artifacts import service

    with pytest.raises(Exception):
        await service.create(
            db,
            workspace_id="agent-1",
            workspace_path=workspace,
            rel_path="../outside.docx",
            data=b"evil",
            created_by="agent-1",
        )
    assert not (Path(workspace).parent / "outside.docx").exists()


async def test_relocate_preserves_artifact_identity(db, workspace, artifact_storage):
    """Rename/move keeps the same Artifact + revision history."""
    from agentos.artifacts import service

    art, rev1 = await service.create(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="drafts/old.docx",
        data=b"v1",
        created_by="agent-1",
    )

    await service.relocate(db, art.id, workspace_path=workspace, new_rel_path="final/report.docx")

    await db.refresh(art)
    assert art.current_path == "final/report.docx"
    assert art.current_revision_id == rev1.id  # same revision, not a new one
    assert (Path(workspace) / "final/report.docx").read_bytes() == b"v1"
    assert not (Path(workspace) / "drafts/old.docx").exists()


async def test_adopt_tracks_an_existing_workspace_file(db, workspace, artifact_storage):
    """Files the user/browser drops into the workspace become tracked."""
    from agentos.artifacts import service

    (Path(workspace) / "download.pdf").write_bytes(b"pdf-v1")

    art, rev = await service.adopt(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="download.pdf",
        created_by="agent-1",
    )

    assert art.format == "pdf"
    assert rev.revision_number == 1
    assert rev.content_hash == hashlib.sha256(b"pdf-v1").hexdigest()
    assert await service.revision_bytes(db, rev.id) == b"pdf-v1"


async def test_archive_tombstones_without_destroying_history(db, workspace, artifact_storage):
    from agentos.artifacts import service

    art, rev1 = await service.create(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="old.docx",
        data=b"v1",
        created_by="agent-1",
    )

    await service.archive(db, art.id)

    await db.refresh(art)
    assert art.archived_at is not None
    # Revisions survive the tombstone — restore/audit stay possible.
    assert await service.revision_bytes(db, rev1.id) == b"v1"


# --- W2b: DOCX structured generation ---


async def test_docx_create_inspect_roundtrip(db, workspace, artifact_storage):
    """Structured spec → valid DOCX → inspectable structure."""
    from agentos.artifacts import service

    spec = {
        "title": "Q3 Report",
        "blocks": [
            {"type": "heading", "level": 1, "text": "Summary"},
            {"type": "paragraph", "text": "Revenue grew 12%."},
            {"type": "list", "style": "bullet", "items": ["alpha", "beta"]},
            {"type": "table", "header": ["Metric", "Value"], "rows": [["ARR", "4.2"]]},
        ],
    }
    art, _rev = await service.create_structured(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="q3.docx",
        format="docx",
        spec=spec,
        created_by="agent-1",
    )

    info = await service.inspect(db, art.id, workspace_path=workspace)
    assert info["valid"] is True
    assert "Summary" in info["structure"]["headings"]
    assert info["structure"]["tables"] == 1
    assert info["structure"]["list_items"] == 2


async def test_docx_revise_applies_ops_as_new_revision(db, workspace, artifact_storage):
    from agentos.artifacts import service

    spec = {"blocks": [{"type": "paragraph", "text": "draft"}]}
    art, rev1 = await service.create_structured(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="memo.docx",
        format="docx",
        spec=spec,
        created_by="agent-1",
    )

    rev2 = await service.revise_structured(
        db,
        art.id,
        workspace_path=workspace,
        base_revision_id=rev1.id,
        ops=[
            {"op": "replace_paragraph", "index": 0, "text": "final"},
            {"op": "append_blocks", "blocks": [{"type": "heading", "level": 1, "text": "Done"}]},
        ],
        change_summary="finalize memo",
    )

    assert rev2.revision_number == 2
    info = await service.inspect(db, art.id, workspace_path=workspace)
    assert "Done" in info["structure"]["headings"]
    # Old revision still inspectable — messages link to exact revisions.
    old = await service.revision_bytes(db, rev1.id)
    assert b"draft" not in old  # zip is binary; verify via structure below
    from agentos.artifacts.formats import docx as docx_fmt

    assert "Done" not in docx_fmt.inspect(old)["structure"]["headings"]


async def test_failed_generation_creates_no_revision(db, workspace, artifact_storage):
    """A build that blows up leaves no artifact rows and no workspace file."""
    from sqlalchemy import func, select

    from agentos.artifacts import service
    from agentos.models.artifact import Artifact, ArtifactRevision

    with pytest.raises(Exception):
        await service.create_structured(
            db,
            workspace_id="agent-1",
            workspace_path=workspace,
            rel_path="bad.docx",
            format="docx",
            spec={"blocks": [{"type": "bogus"}]},
            created_by="agent-1",
        )

    for model in (Artifact, ArtifactRevision):
        n = (await db.execute(select(func.count(model.id)))).scalar_one()
        assert n == 0
    assert not (Path(workspace) / "bad.docx").exists()


async def test_corrupt_file_reports_invalid_not_crash(db, workspace, artifact_storage):
    from agentos.artifacts import service

    # Write garbage where a docx should be, then inspect.
    (Path(workspace) / "corrupt.docx").write_bytes(b"not-a-zip")
    art, _ = await service.adopt(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="corrupt.docx",
        created_by="agent-1",
    )

    info = await service.inspect(db, art.id, workspace_path=workspace)
    assert info["valid"] is False
    assert info["errors"]


# --- W2c: XLSX structured generation ---


async def test_xlsx_create_inspect_with_honest_formula_status(db, workspace, artifact_storage):
    from agentos.artifacts import service

    spec = {
        "sheets": [
            {
                "name": "Summary",
                "rows": [["Metric", "Value"], ["ARR", 4.2], ["MRR", 0.35]],
                "cells": {"B5": {"formula": "=SUM(B2:B3)"}},
                "freeze": "A2",
            }
        ]
    }
    art, _ = await service.create_structured(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="model.xlsx",
        format="xlsx",
        spec=spec,
        created_by="agent-1",
    )

    info = await service.inspect(db, art.id, workspace_path=workspace)
    assert info["valid"] is True
    s = info["structure"]
    assert s["sheets"] == ["Summary"]
    assert s["formulas"] == 1
    # openpyxl writes formulas but never computes — report honestly.
    assert info["formulas_recalculated"] is False


async def test_xlsx_revise_set_cell_and_append(db, workspace, artifact_storage):
    from agentos.artifacts import service

    spec = {"sheets": [{"name": "Data", "rows": [["a", 1]]}]}
    art, rev1 = await service.create_structured(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="book.xlsx",
        format="xlsx",
        spec=spec,
        created_by="agent-1",
    )

    await service.revise_structured(
        db,
        art.id,
        workspace_path=workspace,
        base_revision_id=rev1.id,
        ops=[
            {"op": "set_cell", "sheet": "Data", "cell": "B2", "value": 2},
            {"op": "set_cell", "sheet": "Data", "cell": "B3", "formula": "=SUM(B1:B2)"},
            {"op": "append_rows", "sheet": "Data", "rows": [["b", 3]]},
        ],
    )

    info = await service.inspect(db, art.id, workspace_path=workspace)
    assert info["structure"]["formulas"] == 1
    assert info["structure"]["cells"] == 6  # 2 original + B2 + B3 + appended row


# --- W2d: PPTX structured generation ---


async def test_pptx_create_inspect_roundtrip(db, workspace, artifact_storage):
    from agentos.artifacts import service

    spec = {
        "slides": [
            {"layout": "title", "title": "Q3 Review", "subtitle": "Engineering"},
            {
                "layout": "title_content",
                "title": "Highlights",
                "blocks": [{"type": "bullets", "items": ["ship W1", "start W2"]}],
            },
            {
                "layout": "title_content",
                "title": "Numbers",
                "blocks": [
                    {
                        "type": "chart",
                        "chart_type": "bar",
                        "categories": ["Q1", "Q2", "Q3"],
                        "series": [{"name": "ARR", "values": [3.1, 3.8, 4.2]}],
                    },
                    {"type": "notes", "text": "emphasize Q3 acceleration"},
                ],
            },
        ]
    }
    art, _ = await service.create_structured(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="q3.pptx",
        format="pptx",
        spec=spec,
        created_by="agent-1",
    )

    info = await service.inspect(db, art.id, workspace_path=workspace)
    assert info["valid"] is True
    s = info["structure"]
    assert s["slides"] == 3
    assert s["charts"] == 1


# --- W2e: PDF export + read-only imported PDFs ---


def _pdf_bytes(pages: int = 2) -> bytes:
    from io import BytesIO

    from pypdf import PdfWriter

    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=200, height=200)
    buf = BytesIO()
    w.write(buf)
    return buf.getvalue()


async def test_imported_pdf_is_readable_but_not_editable(db, workspace, artifact_storage):
    from agentos.artifacts import service

    (Path(workspace) / "scan.pdf").write_bytes(_pdf_bytes())
    art, _ = await service.adopt(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="scan.pdf",
        created_by="agent-1",
    )

    info = await service.inspect(db, art.id, workspace_path=workspace)
    assert info["valid"] is True
    assert info["structure"]["pages"] == 2

    # Read-only: no structured revision path for pdf.
    with pytest.raises(service.ArtifactError):
        await service.revise_structured(
            db, art.id, workspace_path=workspace, base_revision_id=art.current_revision_id, ops=[]
        )


async def test_pdf_export_reportlab_fallback_without_soffice(
    db, workspace, artifact_storage, monkeypatch
):
    """No LibreOffice → reportlab renders a valid, tracked PDF."""
    from io import BytesIO

    from pypdf import PdfReader

    from agentos.artifacts import service
    from agentos.models.artifact import Artifact

    monkeypatch.setattr(service, "_find_soffice", lambda: None)

    art, _ = await service.create_structured(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="report.docx",
        format="docx",
        spec={
            "title": "Báo cáo quý",
            "blocks": [
                {"type": "heading", "level": 1, "text": "Tổng quan"},
                {"type": "paragraph", "text": "Chuyên gia phân tích."},
                {"type": "table", "header": ["K", "V"], "rows": [["ok", "1"]]},
            ],
        },
        created_by="agent-1",
    )

    result = await service.export_pdf(db, art.id, workspace_path=workspace)
    assert result["export_status"] == "exported"
    assert result["renderer"] == "reportlab"

    pdf_art = await db.get(Artifact, result["pdf_artifact_id"])
    assert pdf_art.format == "pdf"
    data = await service.revision_bytes(db, pdf_art.current_revision_id)
    reader = PdfReader(BytesIO(data))
    text = reader.pages[0].extract_text()
    assert "Báo cáo quý" in text
    assert "Tổng quan" in text


async def test_pdf_export_unavailable_when_no_renderer_or_extractor(
    db, workspace, artifact_storage, monkeypatch
):
    """renderer_unavailable only when neither soffice nor to_elements exists."""
    from agentos.artifacts import service

    monkeypatch.setattr(service, "_find_soffice", lambda: None)

    art, _ = await service.create(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="data.csv",
        data=b"a,b\n1,2\n",
        created_by="agent-1",
    )

    result = await service.export_pdf(db, art.id, workspace_path=workspace)
    assert result["export_status"] == "renderer_unavailable"
    assert "pdf_artifact_id" not in result


async def test_create_pdf_directly_from_spec(db, workspace, artifact_storage):
    """spec → PDF bytes — pdf is now a creatable (write-once) format."""
    from io import BytesIO

    from pypdf import PdfReader

    from agentos.artifacts import service

    art, rev = await service.create_structured(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="brief.pdf",
        format="pdf",
        spec={
            "title": "Brief",
            "blocks": [{"type": "paragraph", "text": "direct pdf"}],
        },
        created_by="agent-1",
    )

    assert art.format == "pdf"
    data = await service.revision_bytes(db, rev.id)
    reader = PdfReader(BytesIO(data))
    assert "Brief" in reader.pages[0].extract_text()

    # Still write-once: no structured revision path.
    with pytest.raises(service.ArtifactError):
        await service.revise_structured(
            db, art.id, workspace_path=workspace, base_revision_id=rev.id, ops=[]
        )


# --- W2f: mediated capability path ---


def _agent_config(caps: list[str]):
    from agentos.config_schema import AgentConfig, CapabilityGrant, ModelConfig

    return AgentConfig(
        id="art-agent",
        name="Artifact Agent",
        model=ModelConfig(provider_id="test", name="test-model"),
        capabilities=[CapabilityGrant(name=c, require_approval=False) for c in caps],
    )


def _session_stub():
    from types import SimpleNamespace

    return SimpleNamespace(contact_id="c1", id="sess-1", channel=None)


async def test_mediated_artifact_roundtrip_with_provenance(db, workspace, artifact_storage):
    """artifact_create → inspect → revise → history through the mediator,
    with run/message provenance captured on revisions."""
    from sqlalchemy import select

    from agentos.artifacts.service import storage_root  # noqa: F401 — fixture touch
    from agentos.models.artifact import ArtifactRevision
    from agentos.syscall.mediator import SyscallHandler
    from agentos.syscall.protocol import ToolCall

    handler = SyscallHandler(db=db, workspace_path=workspace)
    config = _agent_config(
        ["artifact_create", "artifact_inspect", "artifact_revise", "artifact_history"]
    )
    session = _session_stub()

    spec = {"blocks": [{"type": "heading", "level": 1, "text": "Hi"}]}
    created = await handler.mediate(
        call=ToolCall(
            id="call-1",
            name="artifact_create",
            args={"path": "out.docx", "format": "docx", "spec": spec},
        ),
        session=session,
        agent_config=config,
        run_id="run-art",
    )
    assert created.allowed is True
    art_id = created.output["artifact_id"]
    rev1 = created.output["revision_id"]

    # Provenance captured on the revision row.
    rev_row = (
        await db.execute(select(ArtifactRevision).where(ArtifactRevision.id == rev1))
    ).scalar_one()
    assert rev_row.source_run_id == "run-art"
    assert rev_row.source_message_id == "call-1"

    info = await handler.mediate(
        call=ToolCall(id="call-2", name="artifact_inspect", args={"artifact_id": art_id}),
        session=session,
        agent_config=config,
        run_id="run-art",
    )
    assert info.output["valid"] is True

    revised = await handler.mediate(
        call=ToolCall(
            id="call-3",
            name="artifact_revise",
            args={
                "artifact_id": art_id,
                "base_revision_id": rev1,
                "ops": [{"op": "append_blocks", "blocks": [{"type": "paragraph", "text": "body"}]}],
            },
        ),
        session=session,
        agent_config=config,
        run_id="run-art",
    )
    assert revised.output["revision_number"] == 2

    hist = await handler.mediate(
        call=ToolCall(id="call-4", name="artifact_history", args={"artifact_id": art_id}),
        session=session,
        agent_config=config,
        run_id="run-art",
    )
    assert len(hist.output["revisions"]) == 2
    assert hist.output["revisions"][0]["current"] is True


async def test_ungranted_artifact_call_is_denied(db, workspace, artifact_storage):
    from agentos.syscall.mediator import SyscallHandler
    from agentos.syscall.protocol import ToolCall

    handler = SyscallHandler(db=db, workspace_path=workspace)
    result = await handler.mediate(
        call=ToolCall(
            id="c1",
            name="artifact_create",
            args={"path": "x.docx", "format": "docx", "spec": {}},
        ),
        session=_session_stub(),
        agent_config=_agent_config(["read_file"]),  # no artifact grants
        run_id="run-art",
    )
    assert result.allowed is False


def test_pptx_build_uses_widescreen_canvas():
    """artifact_create decks default to 16:9 — the modern slide standard,
    not python-pptx's dated 4:3 template. Placeholder widths stretch to the
    canvas so content doesn't bunch left."""
    from io import BytesIO

    from pptx import Presentation

    from agentos.artifacts.formats import pptx

    data = pptx.build(
        {"slides": [{"layout": "title_content", "title": "T",
                     "blocks": [{"type": "bullets", "items": ["a"]}]}]},
        ".",
    )
    prs = Presentation(BytesIO(data))
    assert prs.slide_width == 12192000
    assert prs.slide_height == 6858000
    for ph in prs.slides[0].placeholders:
        assert ph.width == prs.slide_width - 2 * ph.left


def test_docx_build_sets_letter_geometry():
    """Page geometry is a deliberate default — US Letter, uniform 1" margins."""
    from io import BytesIO

    import docx

    from agentos.artifacts.formats import docx as docx_fmt

    data = docx_fmt.build({"blocks": [{"type": "paragraph", "text": "hi"}]}, ".")
    section = docx.Document(BytesIO(data)).sections[0]
    assert section.page_width == docx.shared.Inches(8.5)
    assert section.page_height == docx.shared.Inches(11)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        assert getattr(section, m) == docx.shared.Inches(1)


def test_xlsx_build_sets_print_defaults():
    """Sheets paginate sanely: landscape, fit-to-width, modest margins."""
    from io import BytesIO

    from openpyxl import load_workbook

    from agentos.artifacts.formats import xlsx as xlsx_fmt

    data = xlsx_fmt.build({"sheets": [{"name": "S", "rows": [["a"]]}]}, ".")
    ws = load_workbook(BytesIO(data)).active
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_setup.fitToHeight == 0
    assert ws.page_margins.left == 0.5


async def test_create_structured_lands_under_artifacts_dir(db, workspace, artifact_storage):
    """artifact_create paths are rooted at artifacts/ — the deliverables dir,
    mirroring attachments/ for inputs."""
    from agentos.artifacts import service

    spec = {"blocks": [{"type": "paragraph", "text": "hi"}]}
    art, _ = await service.create_structured(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="report.docx",
        format="docx",
        spec=spec,
    )
    assert art.current_path == "artifacts/report.docx"
    assert (Path(workspace) / "artifacts/report.docx").exists()


async def test_create_structured_preserves_subdirs_and_prefix(db, workspace, artifact_storage):
    from agentos.artifacts import service

    spec = {"blocks": [{"type": "paragraph", "text": "hi"}]}
    nested, _ = await service.create_structured(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="reports/q3.docx",
        format="docx",
        spec=spec,
    )
    assert nested.current_path == "artifacts/reports/q3.docx"

    already, _ = await service.create_structured(
        db,
        workspace_id="agent-1",
        workspace_path=workspace,
        rel_path="artifacts/explicit.docx",
        format="docx",
        spec=spec,
    )
    assert already.current_path == "artifacts/explicit.docx"


async def test_create_structured_refuses_attachments_dir(db, workspace, artifact_storage):
    from agentos.artifacts import service

    with pytest.raises(service.ArtifactError, match="attachments"):
        await service.create_structured(
            db,
            workspace_id="agent-1",
            workspace_path=workspace,
            rel_path="attachments/upload.docx",
            format="docx",
            spec={"blocks": [{"type": "paragraph", "text": "hi"}]},
        )
